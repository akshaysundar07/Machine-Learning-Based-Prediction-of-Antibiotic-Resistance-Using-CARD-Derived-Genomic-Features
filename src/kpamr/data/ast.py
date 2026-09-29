"""Table S1 AST processing. No AST label is inferred from genes or MIC values."""

from dataclasses import dataclass
from pathlib import Path
import math
import re

import pandas as pd
from openpyxl import load_workbook

PRIMARY_ANTIBIOTICS = (
    "AMK", "GEN", "TOB", "SAM", "ATM", "CAZ", "CRO", "FEP", "CZA", "C/T",
    "ETP", "IPM", "MEM", "CIP", "LVX", "TZP", "SXT", "TET", "TGC",
)
SOURCE_PATH = Path("data/raw/published_study/mgen-9-967-s002.xlsx")
PHENOTYPE_MAP = {
    "R": "R", "RESISTANT": "R",
    "S": "S", "SUSCEPTIBLE": "S",
    "I": "I", "INTERMEDIATE": "I",
}


def standardize_phenotype(value):
    """Normalize explicit labels; blank cells stay missing, unknowns raise."""
    if pd.isna(value) or (isinstance(value, str) and not value.strip()):
        return pd.NA
    key = str(value).strip().upper()
    if key not in PHENOTYPE_MAP:
        raise ValueError(f"Unexpected AST phenotype: {value!r}")
    return PHENOTYPE_MAP[key]


def add_binary_label(df: pd.DataFrame, phenotype_col: str = "phenotype_std") -> pd.DataFrame:
    out = df.copy()
    phenotypes = out[phenotype_col].map(standardize_phenotype)
    out["label_binary"] = phenotypes.map({"S": 0, "R": 1}).astype("Int64")
    return out


def binary_subset(df: pd.DataFrame, phenotype_col: str = "phenotype_std") -> pd.DataFrame:
    """Return only unambiguous R/S records; preserve I upstream."""
    return df[df[phenotype_col].isin(["R", "S"])].copy()


def read_table_s1(path: str | Path) -> pd.DataFrame:
    """Read the documented A/R:AK schema, stopping at the explicit footnote.

    Validate the entire data block rather than taking the first 100 rows or
    silently dropping nonnumeric IDs. Read formulas literally so an unexpected
    formula cannot be mistaken for a missing cached value.
    """
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "Table S1" not in workbook.sheetnames:
            raise ValueError("Missing worksheet 'Table S1'")
        rows = iter(workbook["Table S1"].iter_rows(max_col=37, values_only=True))
        next(rows, None)  # Published table title.
        header = next(rows, ())
        expected = (*PRIMARY_ANTIBIOTICS, "CST")
        if len(header) != 37 or header[0] != "MRSN IDa" or tuple(header[17:37]) != expected:
            raise ValueError("Table S1 schema mismatch: expected MRSN IDa in A2, "
                             "exactly 19 primary AST columns in R2:AJ2, and CST in AK2")
        if any(header.count(name) != 1 for name in expected):
            raise ValueError("Duplicate AST column in Table S1")
        records = []
        in_footnotes = False
        for row_number, row in enumerate(rows, start=3):
            raw_id = row[0]
            if isinstance(raw_id, str) and raw_id.startswith("Abbreviations -"):
                in_footnotes = True
            if in_footnotes:
                if any(value is not None for value in row[1:]):
                    raise ValueError(f"Unexpected data in Table S1 footnotes at row {row_number}")
                if raw_id is not None and re.fullmatch(r"[0-9]+", str(raw_id).strip()):
                    raise ValueError(f"Isolate after Table S1 footnotes at row {row_number}")
                continue
            isolate_id = str(raw_id).strip() if raw_id is not None else ""
            if not re.fullmatch(r"[0-9]+", isolate_id):
                raise ValueError(f"Invalid/missing MRSN isolate ID at row {row_number}: {raw_id!r}")
            records.append([isolate_id, *row[17:37]])
    finally:
        workbook.close()
    frame = pd.DataFrame(records, columns=["isolate_id", *expected])
    if len(frame) != 100 or frame["isolate_id"].nunique() != 100:
        raise ValueError("Table S1 must contain exactly 100 rows and 100 unique isolate IDs")
    return frame.sort_values("isolate_id", kind="stable").reset_index(drop=True)


def clean_ast_long(wide: pd.DataFrame) -> pd.DataFrame:
    """Retain all source phenotypes, including CST for audit only."""
    required = ["isolate_id", *PRIMARY_ANTIBIOTICS, "CST"]
    if not wide.columns.is_unique or set(required) - set(wide.columns):
        raise ValueError("Missing or duplicate isolate/AST columns")
    ids = wide["isolate_id"]
    if ids.isna().any() or ids.astype(str).str.strip().eq("").any() or ids.duplicated().any():
        raise ValueError("Missing or duplicate isolate IDs")
    tidy = wide[required].melt(
        id_vars="isolate_id", var_name="antibiotic", value_name="original_phenotype"
    )
    normalized = []
    unexpected = []
    for row in tidy.itertuples(index=False):
        try:
            normalized.append(standardize_phenotype(row.original_phenotype))
        except ValueError:
            unexpected.append(f"{row.isolate_id}/{row.antibiotic}={row.original_phenotype!r}")
    if unexpected:
        raise ValueError(f"Unexpected AST labels ({len(unexpected)}): " + "; ".join(unexpected))
    tidy["standardized_phenotype"] = pd.array(normalized, dtype="string")
    tidy = add_binary_label(tidy, "standardized_phenotype").rename(
        columns={"label_binary": "binary_label"}
    )
    tidy["usable_for_binary"] = tidy["binary_label"].notna()
    tidy["primary_antibiotic"] = tidy["antibiotic"].isin(PRIMARY_ANTIBIOTICS)
    return tidy


@dataclass(frozen=True)
class EligibilityRule:
    minimum_usable_binary_count: int = 70
    minimum_minority_class_count: int = 20
    minimum_minority_class_fraction: float = 0.20

    def __post_init__(self):
        for name in ("minimum_usable_binary_count", "minimum_minority_class_count"):
            value = getattr(self, name)
            if type(value) is not int or value < 1:
                raise ValueError(f"{name} must be a positive integer")
        fraction = self.minimum_minority_class_fraction
        if isinstance(fraction, bool) or not isinstance(fraction, (float, int)) or not 0 <= fraction <= 0.5:
            raise ValueError("minimum_minority_class_fraction must be between 0 and 0.5")


def summarize_ast(tidy: pd.DataFrame, rule: EligibilityRule) -> pd.DataFrame:
    """Compute class balance; a zero denominator yields NA, or infinity for n/0."""
    records = []
    for antibiotic, group in tidy.groupby("antibiotic", sort=False):
        phenotype = group["standardized_phenotype"]
        counts = phenotype.value_counts()
        susceptible, intermediate, resistant = (int(counts.get(label, 0)) for label in ("S", "I", "R"))
        missing = int(phenotype.isna().sum())
        if susceptible + intermediate + resistant + missing != len(group):
            raise ValueError(f"Unrecognized standardized phenotype for {antibiotic}")
        usable = susceptible + resistant
        minority = min(susceptible, resistant)
        fraction = minority / usable if usable else math.nan
        ratio = max(susceptible, resistant) / minority if minority else (math.inf if usable else math.nan)
        primary = antibiotic in PRIMARY_ANTIBIOTICS
        reasons = []
        if not primary:
            reasons.append("audit_only_not_primary")
        if usable < rule.minimum_usable_binary_count:
            reasons.append("insufficient_measured_binary_phenotypes")
        if minority < rule.minimum_minority_class_count:
            reasons.append("minority_class_below_minimum")
        if usable == 0 or fraction < rule.minimum_minority_class_fraction:
            reasons.append("minority_fraction_below_minimum")
        records.append({
            "antibiotic": antibiotic,
            "primary_antibiotic": primary,
            "total_isolates": len(group),
            "susceptible_count": susceptible,
            "intermediate_count": intermediate,
            "resistant_count": resistant,
            "missing_count": missing,
            "usable_binary_count": usable,
            "minority_class_count": minority,
            "minority_class_fraction": fraction,
            "majority_to_minority_ratio": ratio,
            "eligible_for_ml": not reasons,
            "eligibility_reason": ";".join(reasons) if reasons else "eligible",
            "minimum_usable_binary_count": rule.minimum_usable_binary_count,
            "minimum_minority_class_count": rule.minimum_minority_class_count,
            "minimum_minority_class_fraction": rule.minimum_minority_class_fraction,
        })
    return pd.DataFrame(records)


def rank_antibiotics(summary: pd.DataFrame) -> pd.DataFrame:
    """Dataset suitability only; alphabetical tie-break makes equal ranks stable."""
    ranked = summary.sort_values(
        ["eligible_for_ml", "minority_class_fraction", "usable_binary_count", "antibiotic"],
        ascending=[False, False, False, True], na_position="last", kind="stable",
    ).reset_index(drop=True)
    ranks = ranked["eligible_for_ml"].cumsum().where(ranked["eligible_for_ml"])
    ranked.insert(0, "initial_experiment_rank", ranks.astype("Int64"))
    return ranked


def binary_label_tables(tidy: pd.DataFrame) -> dict[str, pd.DataFrame]:
    binary = binary_subset(tidy, "standardized_phenotype")
    return {
        antibiotic: binary.loc[
            binary["antibiotic"].eq(antibiotic),
            ["isolate_id", "antibiotic", "standardized_phenotype", "binary_label"],
        ].rename(columns={"standardized_phenotype": "phenotype"}).reset_index(drop=True)
        for antibiotic in PRIMARY_ANTIBIOTICS
    }


def validate_ast_outputs(
    tidy: pd.DataFrame, summary: pd.DataFrame, tables: dict[str, pd.DataFrame],
    expected_isolates: int = 100,
) -> None:
    """Validate identity, mappings, conservation, and exact binary subsets before writing."""
    antibiotics = {*PRIMARY_ANTIBIOTICS, "CST"}
    if set(tidy["antibiotic"]) != antibiotics or set(tables) != set(PRIMARY_ANTIBIOTICS):
        raise ValueError("Expected exactly 19 primary antibiotics and audit-only CST")
    ids = set(tidy["isolate_id"])
    if len(ids) != expected_isolates or tidy.duplicated(["isolate_id", "antibiotic"]).any():
        raise ValueError("Unexpected isolate count or duplicate isolate/antibiotic records")
    expected_labels = tidy["standardized_phenotype"].map({"S": 0, "R": 1}).astype("Int64")
    if not tidy["standardized_phenotype"].dropna().isin(["S", "I", "R"]).all():
        raise ValueError("Unexpected standardized phenotype")
    if not tidy["binary_label"].equals(expected_labels):
        raise ValueError("Invalid S/R/I/missing binary mapping")
    if not tidy["usable_for_binary"].equals(expected_labels.notna()):
        raise ValueError("Invalid binary usability flags")
    if not tidy["primary_antibiotic"].equals(tidy["antibiotic"].isin(PRIMARY_ANTIBIOTICS)):
        raise ValueError("Invalid primary antibiotic flags")
    if summary["antibiotic"].duplicated().any() or set(summary["antibiotic"]) != antibiotics:
        raise ValueError("Invalid summary antibiotic coverage")
    indexed = summary.set_index("antibiotic")
    for antibiotic, group in tidy.groupby("antibiotic", sort=False):
        stats = indexed.loc[antibiotic]
        observed = group["standardized_phenotype"].value_counts()
        if set(group["isolate_id"]) != ids or len(group) != expected_isolates:
            raise ValueError(f"Isolate coverage mismatch for {antibiotic}")
        for label, field in (("S", "susceptible_count"), ("I", "intermediate_count"), ("R", "resistant_count")):
            if stats[field] != observed.get(label, 0):
                raise ValueError(f"Count mismatch for {antibiotic}/{label}")
        if stats["missing_count"] != group["standardized_phenotype"].isna().sum() or stats["total_isolates"] != expected_isolates:
            raise ValueError(f"Missing/total count mismatch for {antibiotic}")
        usable = int(observed.get("S", 0) + observed.get("R", 0))
        if stats["usable_binary_count"] != usable:
            raise ValueError(f"Binary count mismatch for {antibiotic}")
        if antibiotic == "CST":
            if stats["eligible_for_ml"]:
                raise ValueError("CST must remain ineligible")
            continue
        table = tables[antibiotic]
        if len(table) != usable or table["isolate_id"].duplicated().any():
            raise ValueError(f"Invalid binary row count or duplicate ID for {antibiotic}")
        if table["binary_label"].isna().any() or not table["binary_label"].isin([0, 1]).all():
            raise ValueError(f"Invalid binary labels for {antibiotic}")
        expected_table = binary_subset(group, "standardized_phenotype")[
            ["isolate_id", "antibiotic", "standardized_phenotype", "binary_label"]
        ].rename(columns={"standardized_phenotype": "phenotype"}).reset_index(drop=True)
        if not table.equals(expected_table):
            raise ValueError(f"Binary subset mismatch for {antibiotic}")
