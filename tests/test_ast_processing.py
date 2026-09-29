"""M3 regression checks, using synthetic workbooks without modifying raw data."""

import hashlib
import math
from pathlib import Path
import runpy
import sys

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from kpamr.data.ast import (  # noqa: E402
    PRIMARY_ANTIBIOTICS, SOURCE_PATH, EligibilityRule, add_binary_label,
    binary_label_tables, clean_ast_long, rank_antibiotics, read_table_s1,
    standardize_phenotype, summarize_ast, validate_ast_outputs,
)


@pytest.mark.parametrize("raw,expected", [("S", 0), ("R", 1), (" s\t", 0), ("r ", 1)])
def test_binary_mapping(raw, expected):
    result = add_binary_label(pd.DataFrame({"phenotype_std": [raw]}))
    assert result["label_binary"].dtype == "Int64"
    assert result.loc[0, "label_binary"] == expected


@pytest.mark.parametrize("raw", ["I", " i ", None, pd.NA, float("nan"), "", " \t"])
def test_intermediate_and_missing_never_receive_targets(raw):
    assert pd.isna(add_binary_label(pd.DataFrame({"phenotype_std": [raw]})).loc[0, "label_binary"])


@pytest.mark.parametrize("raw", ["NA", "unknown", "ND", "S/R", "-", "=1+1", 0, 1])
def test_unknown_labels_rejected_instead_of_becoming_missing(raw):
    with pytest.raises(ValueError, match="Unexpected AST phenotype"):
        standardize_phenotype(raw)
    with pytest.raises(ValueError, match="Unexpected AST phenotype"):
        add_binary_label(pd.DataFrame({"phenotype_std": [raw]}))


def make_wide(values):
    return pd.DataFrame({"isolate_id": [str(i) for i in range(len(values))],
                         **{antibiotic: list(values) for antibiotic in (*PRIMARY_ANTIBIOTICS, "CST")}})


def test_preservation_counts_and_exact_binary_subsets():
    raw = [" s ", "R", " i\t", None, "", " "]
    tidy = clean_ast_long(make_wide(raw))
    amk = tidy.loc[tidy.antibiotic.eq("AMK")]
    assert amk.original_phenotype.tolist() == raw
    assert amk.standardized_phenotype.iloc[:3].tolist() == ["S", "R", "I"]
    assert amk.standardized_phenotype.iloc[3:].isna().all()
    assert amk.usable_for_binary.tolist() == [True, True, False, False, False, False]
    summary = summarize_ast(tidy, EligibilityRule())
    for row in summary.itertuples():
        assert (row.susceptible_count, row.intermediate_count, row.resistant_count, row.missing_count) == (1, 1, 1, 3)
        assert row.total_isolates == 6
        assert row.usable_binary_count == 2
        assert row.minority_class_fraction == 0.5
        assert row.majority_to_minority_ratio == 1
    tables = binary_label_tables(tidy)
    assert set(tables) == set(PRIMARY_ANTIBIOTICS)
    for antibiotic, table in tables.items():
        assert table.isolate_id.tolist() == ["0", "1"]
        assert table.phenotype.tolist() == ["S", "R"]
        assert table.binary_label.tolist() == [0, 1]
        assert table.antibiotic.eq(antibiotic).all()
    validate_ast_outputs(tidy, summary, tables, expected_isolates=6)
    tables["AMK"].loc[0, "binary_label"] = 1
    with pytest.raises(ValueError, match="Binary subset mismatch"):
        validate_ast_outputs(tidy, summary, tables, expected_isolates=6)


def test_unknown_report_includes_isolate_antibiotic_and_original_value():
    wide = make_wide(["S", "R"])
    wide.loc[1, "C/T"] = "not measured"
    with pytest.raises(ValueError, match="1/C/T='not measured'"):
        clean_ast_long(wide)


@pytest.mark.parametrize("s,r,i,eligible", [(50, 20, 30, True), (80, 20, 0, True),
                                         (81, 19, 0, False), (35, 34, 31, False)])
def test_default_eligibility_boundaries_and_cst_exclusion(s, r, i, eligible):
    summary = summarize_ast(clean_ast_long(make_wide(["S"] * s + ["R"] * r + ["I"] * i)), EligibilityRule())
    assert summary.loc[summary.primary_antibiotic, "eligible_for_ml"].eq(eligible).all()
    assert not summary.loc[summary.antibiotic.eq("CST"), "eligible_for_ml"].item()


def test_configured_fraction_threshold_is_applied():
    tidy = clean_ast_long(make_wide(["S"] * 75 + ["R"] * 25))
    summary = summarize_ast(tidy, EligibilityRule(minimum_minority_class_fraction=0.30))
    assert not summary.eligible_for_ml.any()
    assert summary.eligibility_reason.str.contains("minority_fraction_below_minimum").all()


@pytest.mark.parametrize("values,ratio", [(["S"] * 100, math.inf), ([None] * 100, math.nan)])
def test_zero_class_counts_are_safe(values, ratio):
    tidy = clean_ast_long(make_wide(values))
    summary = summarize_ast(tidy, EligibilityRule())
    assert not summary.eligible_for_ml.any()
    assert summary.minority_class_count.eq(0).all()
    if math.isnan(ratio):
        assert summary.majority_to_minority_ratio.isna().all()
        assert summary.minority_class_fraction.isna().all()
        assert all(table.empty for table in binary_label_tables(tidy).values())
    else:
        assert summary.majority_to_minority_ratio.eq(ratio).all()
        assert summary.minority_class_fraction.eq(0).all()
    validate_ast_outputs(tidy, summary, binary_label_tables(tidy))


def test_ranking_eligibility_fraction_sample_count_then_stable_tie_break():
    summary = pd.DataFrame({
        "antibiotic": ["TGC", "TOB", "GEN", "AMK", "CZA"],
        "eligible_for_ml": [False, True, True, True, True],
        "minority_class_fraction": [0.5, 0.4, 0.4, 0.4, 0.45],
        "usable_binary_count": [60, 100, 90, 100, 80],
    })
    result = rank_antibiotics(summary)
    assert result.antibiotic.tolist() == ["CZA", "AMK", "TOB", "GEN", "TGC"]
    assert result.initial_experiment_rank.iloc[:4].tolist() == [1, 2, 3, 4]
    assert pd.isna(result.initial_experiment_rank.iloc[4])


def make_workbook(path, count=100):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table S1"
    sheet.append(["Table S1. Synthetic regression fixture"])
    sheet.append(["MRSN IDa", *([None] * 16), *PRIMARY_ANTIBIOTICS, "CST"])
    for i in range(count):
        sheet.append([i + 1000, *([None] * 16), *([(" s ", "R", "I", None)[i % 4]] * 19), None])
    sheet.append(["Abbreviations - synthetic table footer"])
    sheet.append(["aIsolate identification number"])
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize("mutation,message", [
    ("missing_column", "schema mismatch"), ("duplicate_column", "schema mismatch"),
    ("missing_id", "Invalid/missing"), ("bad_id", "Invalid/missing"),
    ("duplicate_id", "100 unique"), ("extra_isolate", "100 rows"),
    ("missing_isolate", "100 rows"), ("wrong_sheet", "Missing worksheet"),
    ("footer_data", "footnotes"),
])
def test_workbook_schema_errors_fail_loudly(tmp_path, mutation, message):
    path = tmp_path / "source.xlsx"
    make_workbook(path, count=101 if mutation == "extra_isolate" else 99 if mutation == "missing_isolate" else 100)
    workbook = load_workbook(path)
    sheet = workbook.active
    if mutation == "missing_column":
        sheet["R2"] = None
    elif mutation == "duplicate_column":
        sheet["S2"] = "AMK"
    elif mutation == "missing_id":
        sheet["A3"] = None
    elif mutation == "bad_id":
        sheet["A3"] = "unrecognized row"
    elif mutation == "duplicate_id":
        sheet["A4"] = sheet["A3"].value
    elif mutation == "wrong_sheet":
        sheet.title = "Wrong sheet"
    elif mutation == "footer_data":
        sheet["R103"] = "R"
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match=message):
        read_table_s1(path)


def test_pipeline_serialization_provenance_repeatability_and_unknown_fail_closed(tmp_path):
    source = tmp_path / SOURCE_PATH
    make_workbook(source)
    source_bytes = source.read_bytes()
    (tmp_path / "configs").mkdir()
    (tmp_path / "configs/project.yaml").write_bytes((ROOT / "configs/project.yaml").read_bytes())
    manifest_path = tmp_path / "data/dataset_manifest.csv"
    header = "file_id,local_path,source_type,source_name,source_accession_or_url,retrieved_date,description,sha256,status\n"
    original_row = "SRC001,,bioproject,NCBI,PRJNA717739,,preserved,,PENDING\n"
    manifest_path.write_text(header + original_row, encoding="utf-8")
    process = runpy.run_path(str(ROOT / "scripts/03_clean_ast.py"))["process_ast"]
    outputs = process(tmp_path)
    assert len(outputs) == 22
    assert source.read_bytes() == source_bytes
    assert not (tmp_path / "data/processed/per_antibiotic/CST_labels.csv").exists()
    ct = pd.read_csv(tmp_path / "data/processed/per_antibiotic/CT_labels.csv")
    assert len(ct) == 50
    assert ct.antibiotic.eq("C/T").all()
    assert set(ct.binary_label) == {0, 1}
    assert set(ct.phenotype) == {"S", "R"}
    cleaned = pd.read_csv(tmp_path / "data/interim/ast/ast_clean_long.csv", keep_default_na=False)
    assert " s " in cleaned.original_phenotype.values
    assert cleaned.loc[cleaned.standardized_phenotype.isin(["I", ""]), "binary_label"].eq("").all()
    saved = {p: (tmp_path / p).read_bytes() for p in [*outputs, Path("data/dataset_manifest.csv")]}
    process(tmp_path)
    assert all((tmp_path / p).read_bytes() == content for p, content in saved.items())
    manifest = pd.read_csv(manifest_path, keep_default_na=False)
    assert len(manifest) == 24
    assert manifest.iloc[0].description == "preserved"
    for row in manifest.iloc[1:].itertuples():
        assert hashlib.sha256((tmp_path / row.local_path).read_bytes()).hexdigest() == row.sha256
        if row.source_type == "derived_ast":
            assert row.source_accession_or_url == SOURCE_PATH.as_posix()
    workbook = load_workbook(source)
    workbook.active["R3"] = "NA"  # Literal NA is unknown, not a blank cell.
    workbook.save(source)
    workbook.close()
    with pytest.raises(ValueError, match="1000/AMK='NA'"):
        process(tmp_path)
    assert all((tmp_path / p).read_bytes() == content for p, content in saved.items())
