"""M4 regression tests: reviewed schema, no label leakage, and provenance."""

import hashlib
import json
from pathlib import Path
import runpy
import sys

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from kpamr.features.amr import (  # noqa: E402
    AST_PATH, SOURCE_PATH, GENE_COLUMNS, GENE_HEADERS, NONFEATURE_HEADERS,
    SOURCE_COLUMNS, VARIANT_HEADERS, build_binary_feature_matrix, gene_call,
    machine_feature_id, read_amr_table, remove_invariant_features, summarize_features,
    validate_binary_matrix, validate_m3_isolate_set, variant_tokens,
)


def synthetic_source():
    source = pd.DataFrame(None, index=range(100), columns=["isolate_id", *SOURCE_COLUMNS])
    source["isolate_id"] = [str(1000 + i) for i in range(100)]
    source["AN"] = GENE_HEADERS[0]  # Constant present.
    source.loc[0, "AO"] = "aac(3)-Iid"  # Reviewed case-only spelling.
    source.loc[0, "AP"] = "aac(3)-IIe"  # Distinct rare gene with the same vector.
    source.loc[1, "AT"] = "aac(6')-Ib"
    source.loc[2, "AT"] = "aac(6')-Ib-AKT"
    source.loc[1, "CN"] = "blaSHV-11"
    source.loc[2, "DL"] = "blaSHV-11"
    source.loc[0, "FT"] = "OmpK35-17%;OmpK36GD"
    source.loc[1, "FT"] = "OmpK36GD"
    source.loc[0, "FU"] = "MgrB-0%"
    return source


def feature_id(dictionary, name):
    rows = dictionary.loc[dictionary.display_name.eq(name)]
    assert len(rows) == 1
    return rows.feature_id.item()


@pytest.mark.parametrize("value", [None, pd.NA, float("nan"), "", "  "])
def test_absence_is_zero_without_missing_encoded_values(value):
    source = synthetic_source()
    source["AO"] = value
    full, dictionary = build_binary_feature_matrix(source)
    assert full[feature_id(dictionary, "aac(3)-IId")].eq(0).all()
    assert not full.iloc[:, 1:].isna().any().any()
    assert full.iloc[:, 1:].isin([0, 1]).all().all()


def test_presence_dictionary_constants_and_rare_identical_vectors():
    source = synthetic_source()
    full, dictionary = build_binary_feature_matrix(source)
    model, removed = remove_invariant_features(full)
    assert full.isolate_id.is_unique and len(full) == 100
    assert full.iloc[:, 1:].dtypes.eq("int8").all()
    assert full.columns[1:].tolist() == dictionary.feature_id.tolist()
    assert model.columns[1:].tolist() == dictionary.loc[dictionary.retained_in_model_matrix, "feature_id"].tolist()
    constant_present = feature_id(dictionary, "aac(2')-Iia")
    constant_absent = feature_id(dictionary, "aac(3)-IIg")
    assert full[constant_present].eq(1).all() and constant_present in removed
    assert full[constant_absent].eq(0).all() and constant_absent in removed
    assert dictionary.set_index("feature_id").loc[constant_present, "drop_reason"] == "zero_variance_present"
    assert dictionary.set_index("feature_id").loc[constant_absent, "drop_reason"] == "zero_variance_absent"
    rare = [feature_id(dictionary, name) for name in ("aac(3)-IId", "aac(3)-IIe")]
    assert all(feature in model for feature in rare)
    assert model[rare[0]].equals(model[rare[1]]) and model[rare[0]].sum() == 1
    assert dictionary.feature_id.is_unique
    assert dictionary.feature_id.str.fullmatch(r"[A-Za-z0-9_]+").all()
    iid = dictionary.loc[dictionary.display_name.eq("aac(3)-IId")].iloc[0]
    assert "aac(3)-Iid" in json.loads(iid.published_cell_values)
    assert iid.source_excel_columns == '["AO"]'
    assert json.loads(iid.source_column) == ["aac(3)-IId"]
    summary = summarize_features(source, full, model, dictionary).iloc[0]
    assert summary.zero_variance_present_features == 1
    assert summary.zero_variance_absent_features == len(removed) - 1
    assert summary.identical_vector_groups_retained > 0
    assert summary.missing_encoded_values == summary.duplicate_isolate_ids == 0


def test_duplicate_published_allele_headers_are_or_combined_but_distinct_cell_allele_is_preserved():
    full, dictionary = build_binary_feature_matrix(synthetic_source())
    shv = feature_id(dictionary, "blaSHV-11")
    assert full[shv].sum() == 2
    row = dictionary.set_index("feature_id").loc[shv]
    assert json.loads(row.source_excel_columns) == ["CN", "DL"]
    assert json.loads(row.source_column) == ["blaSHV-11", "blaSHV-11"]
    generic = feature_id(dictionary, "aac(6')-Ib")
    akt = feature_id(dictionary, "aac(6')-Ib-AKT")
    assert full.loc[1, generic] == 1 and full.loc[1, akt] == 0
    assert full.loc[2, generic] == 0 and full.loc[2, akt] == 1
    for name in ("aac(6')-Ib", "aac(6')-Ib-AKT"):
        row = dictionary.loc[dictionary.display_name.eq(name)].iloc[0]
        assert json.loads(row.source_column) == ["aac(6')-Ib-AKT"]
        assert name in json.loads(row.published_names)


@pytest.mark.parametrize("name", [
    "AMK", "C/T", "AMR Phenotypep", "binary_label", "phenotype", "Resistant",
    "Regionb", "Year", "Datec", "Sample Typed", "MLSTe", "All Resistance Genesq",
])
def test_ast_phenotype_metadata_and_summary_columns_cannot_enter_encoding(name):
    source = synthetic_source()
    source[name] = "R"
    with pytest.raises(ValueError, match="reviewed genomic source columns"):
        build_binary_feature_matrix(source)


@pytest.mark.parametrize("value", ["S", "I", "R", "MDR", "XDR", "USA", "Resistant", "NA", "ND", "1", 1, 0, True, "=1+1"])
def test_nonempty_cells_are_not_assumed_to_be_gene_presence(value):
    with pytest.raises(ValueError, match="Unexpected determinant"):
        gene_call("AO", "aac(3)-IId", value)


def test_unreviewed_allele_under_another_header_fails_instead_of_guessing():
    with pytest.raises(ValueError, match="Unexpected determinant"):
        gene_call("AO", "aac(3)-IId", "aac(3)-IIe")


def test_machine_ids_handle_slug_collisions_without_losing_exact_biological_names():
    names = ["aac(6')-Ib", "aac_6_Ib", "fosA7.3", "fosA7-3"]
    ids = [machine_feature_id("gene_or_allele", name) for name in names]
    assert len(set(ids)) == len(names)
    assert all(pd.Series(ids).str.fullmatch(r"[A-Za-z0-9_]+"))


def test_variant_tokens_encode_separately_and_preserve_exact_source_names():
    full, dictionary = build_binary_feature_matrix(synthetic_source())
    omp35 = feature_id(dictionary, "OmpK35-17%")
    omp36 = feature_id(dictionary, "OmpK36GD")
    mgrb = feature_id(dictionary, "MgrB-0%")
    assert full[omp35].sum() == 1
    assert full[omp36].sum() == 2
    assert full[mgrb].sum() == 1  # 0% is a published token, not absence.
    assert full.loc[2, [omp35, omp36, mgrb]].eq(0).all()
    assert variant_tokens("FT", " OmpK36GD ; OmpK35-17%;OmpK36GD ") == ("OmpK35-17%", "OmpK36GD")
    row = dictionary.set_index("feature_id").loc[omp35]
    assert row.display_name == "OmpK35-17%"
    assert json.loads(row.published_cell_values) == ["OmpK35-17%;OmpK36GD"]
    assert row.published_callers == "Kleborate"


@pytest.mark.parametrize("value", ["wild-type", "WT", "OmpK35-100%", "OmpK35-17%,OmpK36GD",
                                  "OmpK35-17%;", "MgrB-0%", "MDR", "=1+1"])
def test_unreviewed_or_wildtype_variant_text_is_never_encoded_as_mutation(value):
    with pytest.raises(ValueError, match="Unexpected variant"):
        variant_tokens("FT", value)


def test_variant_encoding_and_row_order_are_reproducible():
    source = synthetic_source()
    full, dictionary = build_binary_feature_matrix(source)
    shuffled = source.iloc[::-1].copy()
    shuffled.loc[0, "FT"] = "OmpK36GD;OmpK35-17%"
    repeated, changed_dictionary = build_binary_feature_matrix(shuffled)
    pd.testing.assert_frame_equal(full, repeated)
    assert dictionary.feature_id.tolist() == changed_dictionary.feature_id.tolist()


@pytest.mark.parametrize("mutation", ["duplicate_id", "missing_id", "extra_isolate", "fewer_isolates"])
def test_isolate_count_and_identity_errors_are_rejected(mutation):
    source = synthetic_source()
    if mutation == "duplicate_id":
        source.loc[1, "isolate_id"] = source.loc[0, "isolate_id"]
    elif mutation == "missing_id":
        source.loc[0, "isolate_id"] = None
    elif mutation == "extra_isolate":
        source.loc[100] = source.iloc[0]
        source.loc[100, "isolate_id"] = "2000"
    else:
        source = source.iloc[:-1]
    with pytest.raises(ValueError, match="100 unique"):
        build_binary_feature_matrix(source)


def write_ast_ids(path, ids):
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame({"isolate_id": [*ids, *ids], "binary_label": ["not-used"] * (2 * len(ids))}).to_csv(path, index=False)


def test_m3_validation_uses_exact_ids_only(tmp_path):
    source = synthetic_source()
    ast_path = tmp_path / "ast.csv"
    write_ast_ids(ast_path, source.isolate_id)
    validate_m3_isolate_set(source, ast_path)
    altered = source.copy()
    altered.loc[0, "isolate_id"] = "999999"
    with pytest.raises(ValueError, match="M3 isolate set mismatch"):
        validate_m3_isolate_set(altered, ast_path)
    write_ast_ids(ast_path, source.isolate_id.iloc[:-1])
    with pytest.raises(ValueError, match="100 unique"):
        validate_m3_isolate_set(source, ast_path)


@pytest.mark.parametrize("value", [None, 2, "R"])
def test_nonbinary_matrices_cannot_pass_validation(value):
    full, _ = build_binary_feature_matrix(synthetic_source())
    full[full.columns[1]] = value
    with pytest.raises(ValueError, match="binary with no missing"):
        validate_binary_matrix(full)


def make_workbook(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table S1"
    top = [None] * 177
    top[0], top[39], top[175] = "Table S1 synthetic regression fixture", "aac(2')", "Mutationsr"
    sheet.append(top)
    sheet.append([*NONFEATURE_HEADERS, *GENE_HEADERS, *VARIANT_HEADERS])
    for row in synthetic_source().itertuples(index=False, name=None):
        sheet.append([int(row[0]), *(["excluded-source-value"] * 38), *row[1:]])
    sheet.append(["Abbreviations - synthetic source"])
    sheet.append(["qAll resistance genes: AMRFinderPlus and ARIBA; gene name in the row indicates that it is present"])
    sheet.append(["rMutations in ompK35/ompK36 and mgrB as predicted by Kleborate"])
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize("mutation", [
    "gene_header", "ast_as_gene", "variant_header", "extra_column", "missing_column",
    "mutation_family", "missing_gene_note", "missing_variant_note", "missing_boundary",
    "data_after_footer", "bad_id", "duplicate_id", "missing_sheet",
])
def test_workbook_schema_drift_fails_loudly(tmp_path, mutation):
    path = tmp_path / "source.xlsx"
    make_workbook(path)
    workbook = load_workbook(path)
    sheet = workbook.active
    if mutation == "gene_header":
        sheet["AN2"] = "unreviewed_gene"
    elif mutation == "ast_as_gene":
        sheet["AN2"] = "AMK"
    elif mutation == "variant_header":
        sheet["FT2"] = "Resistance phenotype"
    elif mutation == "extra_column":
        sheet["FV2"] = "extra"
    elif mutation == "missing_column":
        sheet.delete_cols(40)
    elif mutation == "mutation_family":
        sheet["FT1"] = "Other"
    elif mutation == "missing_gene_note":
        sheet["A104"] = None
    elif mutation == "missing_variant_note":
        sheet["A105"] = None
    elif mutation == "missing_boundary":
        sheet["A103"] = None
    elif mutation == "data_after_footer":
        sheet["AN103"] = "aac(2')-Iia"
    elif mutation == "bad_id":
        sheet["A3"] = "footnote-like text"
    elif mutation == "duplicate_id":
        sheet["A4"] = sheet["A3"].value
    elif mutation == "missing_sheet":
        sheet.title = "Other"
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError):
        read_amr_table(path)


def test_changes_to_ast_summary_and_metadata_values_cannot_change_features(tmp_path):
    path = tmp_path / "source.xlsx"
    make_workbook(path)
    original, dictionary = build_binary_feature_matrix(read_amr_table(path))
    workbook = load_workbook(path)
    sheet = workbook.active
    for row in sheet.iter_rows(min_row=3, max_row=102, min_col=2, max_col=39):
        for cell in row:
            cell.value = "different-phenotype-or-metadata"
    workbook.save(path)
    workbook.close()
    changed, changed_dictionary = build_binary_feature_matrix(read_amr_table(path))
    pd.testing.assert_frame_equal(original, changed)
    pd.testing.assert_frame_equal(dictionary, changed_dictionary)


def test_pipeline_hashes_protected_inputs_manifest_and_byte_identical_rerun(tmp_path):
    source_path = tmp_path / SOURCE_PATH
    make_workbook(source_path)
    ast_path = tmp_path / AST_PATH
    write_ast_ids(ast_path, synthetic_source().isolate_id)
    (tmp_path / "configs").mkdir()
    config_path = tmp_path / "configs/project.yaml"
    config_path.write_bytes((ROOT / "configs/project.yaml").read_bytes())
    original_inputs = {path: path.read_bytes() for path in (source_path, ast_path, config_path)}
    manifest_path = tmp_path / "data/dataset_manifest.csv"
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    columns = ["file_id", "local_path", "source_type", "source_name", "source_accession_or_url",
               "retrieved_date", "description", "sha256", "status"]
    manifest = pd.DataFrame([
        ["SRC_SOURCE", SOURCE_PATH.as_posix(), "supplementary", "Table S1", "publication",
         "2026-01-01", "preserve source record", source_hash, "VALIDATED"],
        ["M3_KEEP", AST_PATH.as_posix(), "derived_ast", "M3", SOURCE_PATH.as_posix(),
         "", "preserve M3 record", hashlib.sha256(ast_path.read_bytes()).hexdigest(), "VALIDATED_M3"],
    ], columns=columns)
    manifest.to_csv(manifest_path, index=False)
    process = runpy.run_path(str(ROOT / "scripts/04_build_amr_features.py"))["process_amr"]
    outputs = process(tmp_path)
    assert len(outputs) == 4
    assert all(path.read_bytes() == value for path, value in original_inputs.items())
    snapshots = {path: (tmp_path / path).read_bytes() for path in outputs}
    snapshots[Path("data/dataset_manifest.csv")] = manifest_path.read_bytes()
    process(tmp_path)
    assert all((tmp_path / path).read_bytes() == value for path, value in snapshots.items())
    registered = pd.read_csv(manifest_path, dtype=str, keep_default_na=False)
    assert len(registered) == 6 and registered.file_id.is_unique
    pd.testing.assert_frame_equal(registered.iloc[:2].reset_index(drop=True), manifest)
    for row in registered.iloc[2:].itertuples():
        assert row.source_accession_or_url == SOURCE_PATH.as_posix()
        assert row.sha256 == hashlib.sha256((tmp_path / row.local_path).read_bytes()).hexdigest()
    dictionary = pd.read_csv(tmp_path / "data/reports/amr_feature_dictionary.csv", keep_default_na=False)
    full = pd.read_csv(tmp_path / "data/interim/amr/amr_feature_matrix_full.csv", dtype={"isolate_id": str})
    model = pd.read_csv(tmp_path / "data/processed/master/amr_feature_matrix.csv", dtype={"isolate_id": str})
    assert full.shape[0] == model.shape[0] == 100
    assert set(model.columns) <= set(full.columns)
    assert full.columns[1:].tolist() == dictionary.feature_id.tolist()
    assert not full.isna().any().any()
    # Changing outcomes in M3 has no effect on any derived M4 bytes.
    ast = pd.read_csv(ast_path, dtype=str)
    ast["binary_label"] = "another-unused-outcome"
    ast.to_csv(ast_path, index=False)
    process(tmp_path)
    assert all((tmp_path / path).read_bytes() == value for path, value in snapshots.items())
    # A newly unrecognized gene value fails before replacing outputs/manifest.
    workbook = load_workbook(source_path)
    workbook.active["AO3"] = "R"
    workbook.save(source_path)
    workbook.close()
    with pytest.raises(ValueError, match="Unexpected determinant"):
        process(tmp_path)
    assert all((tmp_path / path).read_bytes() == value for path, value in snapshots.items())

